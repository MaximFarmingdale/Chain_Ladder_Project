import chainladder as cl
import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

data = cl.load_sample("clrd")
data = data['CumPaidLoss']
data = data[data["LOB"] == "ppauto"]
state_farm_data = data.loc["State Farm Mut Grp"]
industry_data = data.sum()

print(data.index['GRNAME'].unique())

print(data)
print(industry_data)

sf_dev = cl.Development().fit(state_farm_data)
sf_model = cl.Chainladder().fit(state_farm_data)

industry_dev = cl.Development().fit(industry_data)
industry_model = cl.Chainladder().fit(industry_data)
df_sf_triangle = state_farm_data.to_frame()
df_sf_dev = sf_dev.ldf_.to_frame()
df_ind_ldf = industry_dev.ldf_.to_frame()
df_sf_ult = sf_model.ultimate_.to_frame()
df_sf_ibnr = sf_model.ibnr_.to_frame()

wb = openpyxl.Workbook()
ws_dash = wb.active 
ws_dash.title = "Dashboard"
ws_triangle = wb.create_sheet("Actuarial Triangle")
ws_triangle.cell(row = 1, column = 1, value = "State Farm Cumulative Loss Triangle")
for r_index, row in enumerate(dataframe_to_rows(df_sf_triangle, True, True), start = 3):
    for c_index, value in enumerate(row, start = 1):
        ws_triangle.cell(r_index, c_index, value)
        
        
wb.save("ChainLadder.xlsx")